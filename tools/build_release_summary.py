"""Create the French delivery summary using the Codex bundled Python runtime."""
from pathlib import Path
import argparse, json, statistics
from html import escape

ROOT = Path(__file__).resolve().parents[1]
NUMBERS = [2,3,4,5,8,9,10,11,12,14,15,16,17,18]
FEATURES = [
    (2, 'Déplacer les numéros des fils', 'Maj et glisser sur le numéro'),
    (3, 'Parcourir un circuit', 'Projet / Surligner et parcourir le circuit'),
    (4, 'Numérotation automatique', 'Projet / Numérotation automatique'),
    (5, 'Borniers et raccordement', 'Projet / Borniers et plans de raccordement'),
    (8, 'Implantation d’armoire', 'Projet / Implantation d’armoire'),
    (9, 'Fiches fabricants', 'Projet / Fiches fabricants hors ligne'),
    (10, 'Anomalies localisables', 'Projet / Vérifier les connexions et les repères'),
    (11, 'Recherche globale', 'Édition / Rechercher dans le projet'),
    (12, 'Copies et récupération', 'Fichier / Sauvegarde automatique et récupérer'),
    (14, 'PDF navigable', 'Exporter / Options du PDF navigable'),
    (15, 'Gabarits de projets', 'Fichier / Enregistrer comme gabarit'),
    (16, 'Alignement et répartition', 'Édition / Aligner et répartir la sélection'),
    (17, 'Propriétés et historique plus rapides', 'Automatique dans le panneau Propriétés'),
    (18, 'Présentation commune', 'Affichage / Présentation commune du schéma'),
]
SOURCES = [
    ('Schneider Electric LC1D09BD', 'https://iportal.se.com/Contents/docs/SQD-LC1D09BD.PDF'),
    ('Phoenix Contact PT 2,5 référence 3209510', 'https://www.phoenixcontact.com/en-gb/products/feed-through-terminal-block-pt-25-3209510?type=pdf'),
    ('Pilz PNOZ s4 référence 750104', 'https://www.pilz.com/en-INT/eshop/product/750104'),
    ('Manuel Pilz PNOZ s4 édition EN 14', 'https://www.pilz.com/download/open/PNOZ_s4_Operat_Man_21396-EN-14.pdf'),
]

def content():
    baseline = json.loads((ROOT/'build/benchmark-before-0.18.json').read_text(encoding='utf-8-sig'))['samples']
    measured = json.loads((ROOT/'build/benchmark.json').read_text(encoding='utf-8-sig'))['samples']
    metrics = []
    for old,new in zip(baseline,measured):
        charge = f"{new['components']:,}".replace(',', ' ')
        for label,key in [('Saisie répétée','propertyEditMs'),('Annuler','undoMs'),('Rétablir','redoMs')]:
            a,b = (statistics.median(old[key][1:]),statistics.median(new[key][1:])) if key=='propertyEditMs' else (old[key],new[key])
            metrics.append([charge,label,f'{a:.0f} ms',f'{b:.0f} ms'])
    first = ' et '.join(f"{s['propertyEditMs'][0]:.0f} ms pour {s['components']:,} composants".replace(',', ' ') for s in measured)
    return [
        [
            ('title','GeMMeElec résumé des améliorations'),
            ('meta','Version 0.18 du 13 octobre 2026'),
            ('p','Les 14 améliorations que tu as choisies sont intégrées dans cette version. Ce document explique ce qui change, où trouver chaque commande et quelles limites restent à prendre en compte.'),
            ('p','L’application reste utilisable hors ligne. Les essais ont été réalisés dans des sessions séparées ; le fichier LaserSieve_GH.gemelec a été conservé intact. Les fonctions choisies ont été livrées progressivement en versions 0.16, 0.17 et 0.18.'),
            ('table',['N°','Amélioration','Accès principal'],[[str(n),title,menu] for n,title,menu in FEATURES],[.32,2.30,4.28]),
            ('p','Les numéros correspondent à tes choix. Les menus sont disponibles dans les dispositions Compacte et Ruban. Ctrl+Z annule les changements du projet ; les réglages personnels de l’interface sont mémorisés séparément.'),
        ],
        [
            ('h1','Dessiner et retrouver les objets'),
            ('h2','2 Déplacer les numéros des fils'),
            ('p','Maintenir Maj, cliquer le numéro et glisser pour le placer librement. Le fil conserve ses connexions et son trajet. Les propriétés du fil permettent aussi de régler les décalages, la taille, la couleur, la rotation et la visibilité. Échap abandonne le déplacement ; Ctrl+Z annule le déplacement terminé.'),
            ('h2','3 Parcourir un circuit complet'),
            ('p','Sélectionner un fil puis ouvrir Projet → Surligner / parcourir le circuit. Les fils de sa continuité sont surlignés, y compris à travers les bornes de passage, les pontages déclarés et les reports interfolios. Le tableau permet de rejoindre chaque folio. Effacer le surlignage est accessible depuis le même menu.'),
            ('p','La continuité dépend des ports et liaisons déclarés. Un croisement graphique et une association bobine/contact ne créent pas de connexion électrique. Un numéro indépendant reste inclus dans le parcours du circuit.'),
            ('h2','10 Localiser et actualiser les anomalies'),
            ('p','Projet → Vérifier les connexions et les repères ouvre une liste filtrable par famille, message ou folio. Actualiser relance les contrôles ; Localiser ou un double-clic rejoint l’objet. Les contrôles couvrent notamment les repères manquants ou doublonnés, bornes libres, reports sans destination, références à compléter et numéros incohérents sur une continuité.'),
            ('h2','11 Rechercher dans tout le projet'),
            ('p','Édition → Rechercher dans le projet retrouve les composants, numéros de fils, références, textes et titres de folios. La recherche ignore les accents et la casse. Un double-clic ouvre le bon folio et centre l’objet. Ctrl+F ouvre cette recherche lorsque le dessin a le focus.'),
            ('h2','16 Aligner et répartir les appareils'),
            ('p','Sélectionner plusieurs appareils ou textes, puis Édition → Aligner / répartir la sélection. Les commandes alignent leurs centres sur une ligne ou une colonne et répartissent régulièrement au moins trois objets. Pendant un déplacement, les guides indiquent les centres voisins ; Alt désactive cet accrochage. Les fils suivent les bornes.'),
        ],
        [
            ('h1','Numéroter et documenter les borniers'),
            ('h2','4 Numérotation automatique configurable'),
            ('p','Projet → Numérotation automatique permet de choisir les appareils, les conducteurs ou les deux, pour tout le projet ou un folio. Définir le premier compteur, le préfixe des fils et les formats. Exemples : {prefix}{n:000} donne K001 ; {prefix}{folio}-{n:00} donne W3-01. Le compteur peut repartir par folio.'),
            ('p','Actualiser l’aperçu affiche les anciennes et nouvelles valeurs. Les conflits empêchent l’application ; une modification des règles impose de revoir l’aperçu. Appliquer réalise les changements en une action annulable. Les règles sont enregistrées avec le projet.'),
            ('p','La case de verrouillage du repère, dans les propriétés d’un appareil, protège sa famille bobine/contact. Les fils marqués Numéro indépendant conservent leur numéro. Un réseau électrique qui traverse plusieurs folios garde un numéro commun, même lorsque le périmètre demandé est un seul folio.'),
            ('h2','5 Borniers structurés et plans de raccordement'),
            ('p','Placer puis sélectionner les symboles de bornes sur le schéma. Dans Projet → Borniers et plans de raccordement, créer un bornier, lui donner un nom et une localisation, puis Ajouter la sélection du schéma.'),
            ('p','Chaque affectation comporte un numéro, un niveau, les ports interne et externe et un commentaire. L’onglet Pontages relie des ports existants et participe réellement à la continuité électrique. Les plans affichent les départs, arrivées, numéros de fils et adresses folio/zone.'),
            ('p','Enregistrer le bornier conserve ses affectations. Générer les folios du plan crée ou régénère les pages dans le projet ; les longs commentaires et les pontages sont paginés. Les pages font partie du PDF et ne se modifient pas directement sur le dessin.'),
            ('p','Un niveau supplémentaire exige une paire de ports avec une continuité explicitement définie dans le symbole. On peut aussi affecter deux symboles distincts au même numéro avec des niveaux différents. Les deux ports d’une bobine ne constituent pas une borne de passage.'),
        ],
        [
            ('h1','Implanter une armoire et consulter les fiches'),
            ('h2','8 Implantation d’armoire'),
            ('p','Projet → Implantation d’armoire ouvre une vue de face en millimètres. Créer l’enveloppe, puis ajouter les appareils sélectionnés du schéma. Leurs dimensions sont saisies explicitement ou reprises d’une fiche exacte lorsque la largeur et la hauteur sont connues.'),
            ('p','Ajouter des rails et goulottes, régler les dimensions et les positions dans le tableau, puis glisser les objets dans l’aperçu. Le déplacement s’accroche tous les 5 mm. Enregistrer conserve l’implantation ; Générer le folio ajoute la vue et les pages de cotes au projet.'),
            ('p','Vérifier les chevauchements signale les objets qui se recouvrent et les appareils physiques représentés plusieurs fois. Le placement sur un rail est autorisé. Les objets sont des encombrements rectangulaires en vue de face ; les dégagements thermiques, profondeurs utiles et règles de montage fabricant restent à vérifier.'),
            ('h2','9 Fiches fabricants disponibles hors ligne'),
            ('p','Projet → Fiches fabricants hors ligne permet de chercher une référence, lire ses caractéristiques et affecter les données au composant sélectionné. La fiche conserve sa source et son édition. L’affectation ne change pas la géométrie ni le bornage électrique du symbole.'),
            ('table',['Référence exacte','Disponibilité et édition'],[
                ['Schneider LC1D09BD','Fiche et PDF embarqué du 8 septembre 2017. Cotes issues de cette édition.'],
                ['Phoenix Contact 3209510 PT 2,5','Caractéristiques hors ligne vérifiées le 13 octobre 2026. PDF non embarqué car son serveur a refusé le téléchargement.'],
                ['Pilz 750104 PNOZ s4','Fiche exacte et manuel PDF EN 14 embarqué. Largeur documentée ; hauteur et profondeur non renseignées.'],
            ],[2.25,4.65]),
            ('p','Ces trois fiches constituent une sélection limitée. Les éditions anciennes doivent être confrontées au matériel livré. Aucune caractéristique ni dimension manquante n’a été déduite. Les bibliothèques et les documents embarqués restent consultables sans Internet.'),
            ('links',SOURCES),
        ],
        [
            ('h1','Sauvegarder réutiliser et exporter'),
            ('h2','12 Sauvegarde automatique et récupération'),
            ('p','Fichier → Sauvegarde automatique / récupérer règle l’activation, l’intervalle et le nombre de copies. Valeurs par défaut : une copie toutes les 120 secondes pour un projet modifié, avec cinq copies par projet. L’intervalle va de 30 à 3 600 secondes et la conservation de 1 à 20 copies.'),
            ('p','Les copies sont compressées dans le dossier local GeMMeElec-recovery du compte Windows. Une écriture s’exécute en arrière-plan à partir d’un instantané, hors déplacement en cours. Les copies disponibles sont proposées au démarrage et peuvent être ouvertes depuis le tableau de récupération.'),
            ('p','Restaurer ouvre un projet non enregistré : utiliser Enregistrer sous pour garder le fichier d’origine. L’intégrité de la copie est contrôlée avant ouverture. La copie automatique complète l’enregistrement manuel ; elle ne récupère pas les changements intervenus depuis sa dernière exécution.'),
            ('h2','14 PDF navigable'),
            ('p','Exporter → Options du PDF navigable propose trois réglages : sommaire avec liens vers les folios, signets de navigation et liens sur les reports ou références bobine/contact. Exporter ensuite Tous les folios en PDF. Le sommaire est paginé et ne crée pas de folio dans le projet.'),
            ('p','Les nouveaux projets activent ces options. À l’ouverture d’un ancien projet, le sommaire reste désactivé pour conserver le nombre de pages habituel ; il peut être activé manuellement. Les numéros de folios ne changent pas quand le PDF reçoit un sommaire.'),
            ('h2','15 Gabarits de projets'),
            ('p','Fichier → Enregistrer comme gabarit crée un fichier .gemtemplate avec les cartouches, images, symboles personnels, repérage, règles et organisation des folios. Par défaut, le nouveau dossier contient des schémas vides. Une case permet d’inclure les appareils, fils, borniers et implantations ; une autre inclut la palette et la couleur des prochains fils.'),
            ('p','Fichier → Nouveau depuis un gabarit demande le nom du dossier et crée des identifiants indépendants, avec les associations remappées. La palette incluse est reprise dans les préférences de dessin. Le projet créé doit ensuite être enregistré en .gemelec ; le gabarit reste disponible pour d’autres dossiers.'),
        ],
        [
            ('h1','Présentation performances et vérifications'),
            ('h2','18 Présentation commune avec aperçu'),
            ('p','Affichage → Présentation commune du schéma propose Technique, Noir et Lecture agrandie. Ajuster tailles des repères, descriptions, caractéristiques et numéros de fils, puis couleurs des symboles et fils et épaisseur commune. Actualiser l’aperçu montre le résultat sur une copie du folio avant application.'),
            ('p','Conserver les réglages individuels protège les styles personnalisés et les conducteurs à numéro indépendant. Décocher autorise l’uniformisation. Une modification après aperçu demande de vérifier le nouvel aperçu avant de cliquer de nouveau sur Appliquer. Les images fabricants restent dans leurs couleurs d’origine ; les positions et connexions sont conservées.'),
            ('h2','17 Propriétés et annulation plus rapides'),
            ('p','La saisie du repère, de la description, des données d’article, des coordonnées, de la rotation et des libellés de bornes utilise des transactions limitées aux folios concernés. Les contacts associés restent synchronisés et l’historique local se combine avec les opérations structurelles.'),
            ('table',['Composants','Opération','Version 0.17','Version 0.18'],metrics,[1.1,2.25,1.75,1.8]),
            ('p','Mesures locales en processus isolés, sur 100 et 200 folios avec 4 900 et 9 800 fils. La saisie répétée est la médiane des quatre modifications après la première. La première modification mesurée reste à '+first+'. Les opérations globales et certains exports demeurent coûteux ; ces résultats ne mesurent pas une cadence GPU ou de souris.'),
            ('h2','Vérifications et compatibilité'),
            ('p','129 scénarios applicatifs et 118 contrôles MCP ont réussi sur la livraison, avec 68 outils MCP. Sauvegarde, relecture, connexions, annulation et migrations des formats 1 à 6 vers le format 7 sont vérifiées. Les liens et signets du PDF ont été relus puis les pages rendues et inspectées.'),
            ('p','Limites conservées : 200 folios, fichier de projet de 64 Mo, contrôles de cohérence sans certification électrique, import QET partiel et DXF graphique. Les anciens exécutables ne lisent pas le format 7. Les plans générés se régénèrent depuis leur gestionnaire après modification de leurs données.'),
        ],
    ]

def make_docx(pages):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    doc=Document()
    for style in doc.styles:
        for border in style.element.xpath('./w:pPr/w:pBdr'):
            border.getparent().remove(border)
    sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
    sec.header_distance=Inches(.25);sec.footer_distance=Inches(.3)
    for name,size in [('Normal',11),('Title',24),('Subtitle',11),('Heading 1',18),('Heading 2',12)]:
        style=doc.styles[name];style.font.name='Calibri';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
        style.paragraph_format.space_after=Pt(7);style.paragraph_format.line_spacing=1.05
        if name in ('Heading 1','Heading 2'):style.paragraph_format.space_before=Pt(10);style.paragraph_format.keep_with_next=True
    doc.styles['Title'].paragraph_format.space_after=Pt(8)
    head=sec.header.paragraphs[0];head.text='GeMMeElec  |  Modifications sélectionnées';head.style='Normal';head.runs[0].font.size=Pt(9)
    foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    foot.add_run('Version 0.18  •  13 octobre 2026  •  Page ')
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
    for run in foot.runs:run.font.size=Pt(9)
    for page_index,page in enumerate(pages):
        if page_index:doc.add_page_break()
        for item in page:
            kind=item[0]
            if kind in ('title','meta','h1','h2','p'):
                style={'title':'Title','meta':'Subtitle','h1':'Heading 1','h2':'Heading 2','p':'Normal'}[kind];p=doc.add_paragraph(item[1],style)
                if kind=='meta':p.paragraph_format.space_after=Pt(13)
            elif kind=='table':
                _,headers,rows,widths=item;table=doc.add_table(rows=1,cols=len(headers));table.autofit=False
                for i,w in enumerate(widths):table.columns[i].width=Inches(w)
                for row,values in zip([table.rows[0]]+[table.add_row() for _ in rows],[headers]+rows):
                    for i,(cell,text) in enumerate(zip(row.cells,values)):
                        cell.width=Inches(widths[i]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        cell.text=str(text);tcpr=cell._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
                        for edge,value in [('top','75'),('bottom','75'),('left','100'),('right','100')]:e=OxmlElement('w:'+edge);e.set(qn('w:w'),value);e.set(qn('w:type'),'dxa');marg.append(e)
                        tcpr.append(marg)
                        for p in cell.paragraphs:
                            p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0
                            if i==0 and headers[0]=='N°':p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                            for run in p.runs:run.font.size=Pt(10)
                for c in table.rows[0].cells:
                    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EAF0F5');c._tc.get_or_add_tcPr().append(sh)
                    for run in c.paragraphs[0].runs:run.bold=True
                repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
                borders=OxmlElement('w:tblBorders')
                for edge in ['top','left','bottom','right','insideH','insideV']:
                    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
                table._tbl.tblPr.append(borders)
                doc.add_paragraph().paragraph_format.space_after=Pt(0)
            elif kind=='links':
                for label,url in item[1]:
                    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3)
                    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');prop=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'175B8C');prop.append(color);size=OxmlElement('w:sz');size.set(qn('w:val'),'18');prop.append(size);r.append(prop);t=OxmlElement('w:t');t.text=label;r.append(t);link.append(r);p._p.append(link)
    path=ROOT/'output/documents/GeMMeElec_Modifications_0.18.docx';path.parent.mkdir(parents=True,exist_ok=True)
    doc.core_properties.title='GeMMeElec résumé des améliorations';doc.core_properties.subject='Modifications sélectionnées du 13 octobre 2026';doc.core_properties.author='GeMMeElec';doc.save(path)
    print(path)

def make_pdf(pages):
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer, PageBreak
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    fonts=Path('C:/Windows/Fonts');pdfmetrics.registerFont(TTFont('Calibri',str(fonts/'calibri.ttf')));pdfmetrics.registerFont(TTFont('Calibri-Bold',str(fonts/'calibrib.ttf')));pdfmetrics.registerFontFamily('Calibri',normal='Calibri',bold='Calibri-Bold')
    styles={}
    definitions={'p':(11,13.4,0,7),'title':(24,28,0,8),'meta':(11,13,0,13),'h1':(18,22,0,13),'h2':(12,15,9,6),'cell':(10,12,0,0),'link':(9,11,0,3)}
    for name,(size,leading,before,after) in definitions.items():styles[name]=ParagraphStyle(name,fontName='Calibri-Bold' if name in ('h1','h2') else 'Calibri',fontSize=size,leading=leading,spaceBefore=before,spaceAfter=after,textColor=colors.black,keepWithNext=name in ('h1','h2'))
    path=ROOT/'output/pdf/GeMMeElec_Modifications_0.18.pdf';path.parent.mkdir(parents=True,exist_ok=True)
    document=SimpleDocTemplate(str(path),pagesize=letter,rightMargin=57.6,leftMargin=57.6,topMargin=46.8,bottomMargin=46.8,title='GeMMeElec résumé des améliorations',author='GeMMeElec')
    flow=[]
    for index,page in enumerate(pages):
        if index:flow.append(PageBreak())
        for item in page:
            kind=item[0]
            if kind in definitions:flow.append(Paragraph(escape(item[1]),styles[kind]))
            elif kind=='table':
                _,headers,rows,widths=item;values=[[Paragraph(escape(str(x)),styles['cell']) for x in row] for row in [headers]+rows];table=Table(values,colWidths=[w*72 for w in widths],repeatRows=1,hAlign='LEFT');table.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.4,colors.HexColor('#D9D9D9')),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#EAF0F5')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5),('TOPPADDING',(0,0),(-1,-1),4),('BOTTOMPADDING',(0,0),(-1,-1),4)]));flow.extend([table,Spacer(1,7)])
            elif kind=='links':
                for label,url in item[1]:flow.append(Paragraph(f'<link href="{escape(url,quote=True)}" color="#175B8C">{escape(label)}</link>',styles['link']))
    def frame(canvas,doc):
        canvas.setFont('Calibri',9);canvas.setFillColor(colors.black);canvas.drawString(57.6,766,'GeMMeElec  |  Modifications sélectionnées');canvas.drawRightString(554.4,25,f'Version 0.18  •  13 octobre 2026  •  Page {doc.page}')
    document.build(flow,onFirstPage=frame,onLaterPages=frame);print(path)

def make_markdown(pages):
    result=[]
    for page in pages:
        for item in page:
            if item[0] in ('title','h1','h2'):result.append(('#' if item[0]=='title' else '##' if item[0]=='h1' else '###')+' '+item[1]+'\n')
            elif item[0] in ('p','meta'):result.append(item[1]+'\n')
            elif item[0]=='table':
                result+=['| '+' | '.join(item[1])+' |','| '+' | '.join(['---']*len(item[1]))+' |'];result+=['| '+' | '.join(map(str,row))+' |' for row in item[2]];result.append('')
            elif item[0]=='links':result+=['- ['+label+']('+url+')' for label,url in item[1]]+['']
    (ROOT/'docs/LIVRAISON_0.18.md').write_text('\n'.join(result),encoding='utf-8')

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--format',choices=['docx','pdf','markdown'],required=True);args=parser.parse_args();pages=content()
    {'docx':make_docx,'pdf':make_pdf,'markdown':make_markdown}[args.format](pages)
